"""Reading rosters from Excel and writing the finished schedule back to Excel."""

from __future__ import annotations

import io
from typing import List

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from scheduler import Player, Schedule, ScheduleError

# Column names we will accept (case-insensitive).
NAME_ALIASES = {"name", "player", "player name", "full name"}
FIRST_NAME_ALIASES = {"first name", "first", "firstname"}
LAST_NAME_ALIASES = {"last name", "last", "lastname", "surname"}
LEVEL_ALIASES = {"level", "usda", "usda level", "ntrp", "ntrp level", "rating",
                 "usta", "usta level", "tournament rating"}


def read_roster(source) -> List[Player]:
    """
    Read a roster from an Excel file (path or file-like object).

    Player names come from a single name column, or from separate First/Last
    name columns. The level comes from a level column (Level, NTRP, USDA,
    Tournament Rating, ...). Headers are matched flexibly.
    """
    df = pd.read_excel(source)
    if df.empty:
        raise ScheduleError("The spreadsheet has no rows.")

    lookup = {str(c).strip().lower(): c for c in df.columns}
    name_col = next((lookup[k] for k in lookup if k in NAME_ALIASES), None)
    first_col = next((lookup[k] for k in lookup if k in FIRST_NAME_ALIASES), None)
    last_col = next((lookup[k] for k in lookup if k in LAST_NAME_ALIASES), None)
    level_col = next((lookup[k] for k in lookup if k in LEVEL_ALIASES), None)

    has_name = name_col is not None or (first_col is not None or last_col is not None)
    if not has_name or level_col is None:
        raise ScheduleError(
            "Could not find the required columns. Expected a name column "
            f"(one of {sorted(NAME_ALIASES)}) or First/Last name columns, and a level column "
            f"(one of {sorted(LEVEL_ALIASES)}). Found columns: {list(df.columns)}."
        )

    def row_name(row) -> str:
        if name_col is not None and not pd.isna(row[name_col]):
            return str(row[name_col]).strip()
        parts = []
        for col in (first_col, last_col):
            if col is not None and not pd.isna(row[col]):
                parts.append(str(row[col]).strip())
        return " ".join(parts).strip()

    players: List[Player] = []
    missing_level: List[str] = []
    for _, row in df.iterrows():
        nm = row_name(row)
        raw_level = row[level_col]
        if nm == "":
            continue  # skip blank rows
        if pd.isna(raw_level) or str(raw_level).strip() == "":
            missing_level.append(nm)
            continue
        try:
            level = float(raw_level)
        except (TypeError, ValueError):
            raise ScheduleError(
                f"Player '{nm}' has a non-numeric level: '{raw_level}'. "
                "Levels must be numbers like 3.0 or 4.5."
            )
        players.append(Player(name=nm, level=level))

    if missing_level:
        raise ScheduleError(
            "These players have no level/rating and cannot be scheduled: "
            f"{', '.join(missing_level)}. Add a rating for each, or remove them from the file."
        )

    if not players:
        raise ScheduleError("No valid players found in the spreadsheet.")

    names = [p.name for p in players]
    dupes = {x for x in names if names.count(x) > 1}
    if dupes:
        raise ScheduleError(f"Duplicate player names found: {sorted(dupes)}. Names must be unique.")

    return players


# ---------------------------------------------------------------------------
# Excel writing
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_ROUND_FILL = PatternFill("solid", fgColor="2E7D32")
_ROUND_FONT = Font(color="FFFFFF", bold=True, size=12)
_THIN = Side(style="thin", color="D0D0D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")


def _style_header(cell) -> None:
    cell.fill = _HEADER_FILL
    cell.font = _HEADER_FONT
    cell.alignment = _CENTER
    cell.border = _BORDER


def _autosize(ws) -> None:
    for col_cells in ws.columns:
        width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(width + 3, 40)


def build_workbook(schedule: Schedule) -> Workbook:
    wb = Workbook()
    _write_schedule_sheet(wb.active, schedule)
    _write_player_sheet(wb.create_sheet("By Player"), schedule)
    return wb


def _write_schedule_sheet(ws, schedule: Schedule) -> None:
    ws.title = "Schedule"
    headers = ["Court", "Team A", "Level", "vs", "Team B", "Level",
               "Court Avg", "Level Spread", "Team Gap"]
    row = 1
    for rnd in schedule.rounds:
        # Round banner
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        banner = ws.cell(row=row, column=1, value=f"ROUND {rnd.number}")
        banner.fill = _ROUND_FILL
        banner.font = _ROUND_FONT
        banner.alignment = _LEFT
        row += 1

        for c, h in enumerate(headers, start=1):
            _style_header(ws.cell(row=row, column=c, value=h))
        row += 1

        for m in sorted(rnd.matches, key=lambda x: x.court):
            a1, a2 = (schedule.players[i].name for i in m.team_a)
            b1, b2 = (schedule.players[i].name for i in m.team_b)
            avg = sum(schedule.players[i].level for i in m.all_indices()) / 4
            values = [
                m.court,
                f"{a1} & {a2}",
                round(schedule.team_level(m.team_a) / 2, 2),
                "vs",
                f"{b1} & {b2}",
                round(schedule.team_level(m.team_b) / 2, 2),
                round(avg, 2),
                round(schedule.court_spread(m), 2),
                round(schedule.team_gap(m), 2),
            ]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.border = _BORDER
                cell.alignment = _CENTER if c != 2 and c != 5 else _LEFT
            row += 1

        if rnd.byes:
            bye_names = ", ".join(schedule.players[i].name for i in rnd.byes)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
            cell = ws.cell(row=row, column=1, value=f"Byes: {bye_names}")
            cell.font = Font(italic=True, color="8A6D3B")
            cell.alignment = _LEFT
            row += 1

        row += 1  # blank spacer between rounds

    _autosize(ws)
    ws.freeze_panes = "A2"


def _write_player_sheet(ws, schedule: Schedule) -> None:
    headers = ["Player", "Level", "Matches", "Byes", "Partners", "Opponents"]
    for c, h in enumerate(headers, start=1):
        _style_header(ws.cell(row=1, column=c, value=h))

    for r, s in enumerate(sorted(schedule.player_stats(), key=lambda x: (-x["level"], x["name"])), start=2):
        values = [
            s["name"], s["level"], s["matches"], s["byes"],
            ", ".join(s["partners"]), ", ".join(s["opponents"]),
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = _BORDER
            cell.alignment = _CENTER if c in (2, 3, 4) else _LEFT

    _autosize(ws)
    ws.freeze_panes = "A2"


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
